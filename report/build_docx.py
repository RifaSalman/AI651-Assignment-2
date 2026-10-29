"""Build an editable Word report with real tables."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "Question2" / "kaggle_output" / "outputs"
if not (FIG / "baseline_scores.json").exists():
    FIG = ROOT / "Question2" / "outputs"
OUT = ROOT / "report" / "AI651_Assignment2_RifaSalman.docx"
REPO = "https://github.com/RifaSalman/AI651-Assignment-2"

SCORES = json.loads((FIG / "baseline_scores.json").read_text(encoding="utf-8"))
PHASE2 = json.loads((FIG / "phase2_scores.json").read_text(encoding="utf-8"))


def fmt(value: float) -> str:
    if value != value:
        return "—"
    magnitude = abs(value)
    if magnitude == 0:
        return "0"
    if magnitude < 1e-3:
        text = f"{value:.2e}"
        return text.replace("e-0", "e-").replace("e+0", "e+")
    return f"{value:.3g}"


def pair(a: float, b: float) -> str:
    return f"{fmt((a + b) / 2)} ({fmt(a)}, {fmt(b)})"


def learned_row(block: list[dict], key: str) -> str:
    return pair(block[0][key], block[1][key])


def cell_mean(rows: list[dict], *keys: str) -> str:
    values = []
    for row in rows:
        item = row
        for key in keys:
            item = item[key]
        values.append(float(item))
    return pair(values[0], values[1])


def set_run_font(run, size: float = 11, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def style_paragraph(para, space_after: float = 8) -> None:
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.line_spacing = 1.15


def add_body(doc: Document, text: str, *, size: float = 11, bold: bool = False, italic: bool = False) -> None:
    para = doc.add_paragraph()
    style_paragraph(para)
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)


def add_caption(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    style_paragraph(para, space_after=10)
    run = para.add_run(text)
    set_run_font(run, size=9, italic=True)


def add_heading_text(doc: Document, text: str, level: int) -> None:
    para = doc.add_heading(text, level=level)
    for run in para.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True)


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tePr if False else cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 8) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], font: float = 8) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, size=font)
        shade_cell(table.rows[0].cells[index], "D9E2F3")
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            set_cell_text(table.rows[r_index + 1].cells[c_index], value, size=font)
    doc.add_paragraph()


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.2) -> None:
    if path.exists():
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(para, space_after=4)
        para.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def render_q1(doc: Document) -> None:
    text = (ROOT / "report" / "q1.md").read_text(encoding="utf-8")
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            add_heading_text(doc, line[2:].strip(), 1)
        elif line.startswith("## "):
            add_heading_text(doc, line[3:].strip(), 2)
        elif line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            cells = []
            for raw in block:
                if set(raw.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
                    continue
                cells.append([part.strip() for part in raw.strip("|").split("|")])
            if cells:
                add_table(doc, cells[0], cells[1:], font=8)
            continue
        elif line.startswith("- "):
            para = doc.add_paragraph(style="List Bullet")
            style_paragraph(para)
            run = para.add_run(line[2:].strip())
            set_run_font(run, size=11)
        else:
            para_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "- ")):
                para_lines.append(lines[i].rstrip())
                i += 1
            add_body(doc, " ".join(para_lines))
            continue
        i += 1
    add_figure(
        doc,
        ROOT / "report" / "figures" / "q1_architecture.png",
        "Figure Q1. Forecasting architecture. Blue boxes are observations, sand boxes are fixed by the river, "
        "green boxes are learned, and the issued product is the +60 min level at S1–S4.",
        width=6.4,
    )


def add_q2(doc: Document) -> None:
    a = SCORES["A"]
    blank = SCORES["reference"]["Blank"]
    copy = SCORES["reference"]["Copy-last"]
    orig = a["original"]
    models = PHASE2["models"]
    wrapped = PHASE2["wrapped"]

    add_heading_text(doc, "Question 2 — Motion or Memory?", 1)
    add_body(
        doc,
        "Hardware: NVIDIA Tesla T4 (Kaggle). Training time is mean wall-clock seconds per epoch over 30 epochs "
        "(dataset download excluded). Inference time is wall-clock time for all 1,000 original test videos at batch 32, "
        "after each batch is on device, under eval() and no_grad(). "
        "Learned entries are mean (seed 0, seed 1). Figures and GIFs use seed 0.",
    )

    add_heading_text(doc, "Table 1. Transformation and window checks", 2)
    add_table(
        doc,
        ["Check", "Expected", "Result"],
        [
            ["(a) max |M(M(x)) − x|", "0", "0, torch.equal True"],
            ["(b) max |R^4(x) − x|", "0", "0, torch.equal True"],
            ["(c) max |M(R(M(x))) − R^{−1}(x)|", "0", "0, torch.equal True"],
            ["(d) bright pixel of M(p)", "(7, 10, 43)", "(7, 10, 43)"],
            ["(e) bright pixel of R(p)", "(7, 43, 10)", "(7, 43, 10)"],
            ["(f) 1× input / target", "1–5 / 6–10", "1–5 / 6–10"],
            ["(g) 2× input / target", "1, 3, 5, 7, 9 / 11, 13, 15, 17, 19", "same"],
            ["(h) 3× input / target", "1, 4, 7, 10, 13 / 16, 19", "same"],
        ],
    )
    add_caption(
        doc,
        "Table 1. Checks (a)–(h) on the functions used in the experiments. "
        "x is torch.rand(20, 1, 64, 64) after torch.manual_seed(0). "
        "p is zero except p[7, 0, 10, 20] = 1. Model A has 480,784 parameters.",
    )

    stats = json.loads((FIG / "figure1_stats.json").read_text(encoding="utf-8"))
    add_figure(
        doc,
        ROOT / "Question2" / "outputs" / "figure1.png",
        "Figure 1. Optical flow on the first 50 training videos, 950 frame pairs, Farneback settings as specified. "
        "Pixels kept: intensity above 0.2 on both frames and magnitude above 0.5. "
        f"Kept pixels: {stats['n_pixels']}. Median speed 3.32 pixels per frame. The horizontal line in (a) is 1/18.",
    )
    add_heading_text(doc, "Q2.1", 2)
    add_body(
        doc,
        "Mirrored and rotated trajectories are partly familiar, and they are not exchangeable with the training distribution. "
        "Every 20° bin in Figure 1(a) is occupied, so a reflected or rotated direction is not a new kind of motion. "
        "The bins are not flat: the fullest holds 0.082 of the kept pixels and the emptiest 0.026, a ratio of 3.11, against a uniform level of 1/18 ≈ 0.0556. "
        "Reflecting every direction moves about 0.13 of the probability mass (total variation between the histogram and its mirror). "
        "A 90° rotation is not aligned to the 20° bins; the closest shifts, 4 and 5 bins, move about 0.18 of the mass. "
        "A typical speed is the median, 3.32 pixels per frame (mean 3.27, 90th percentile 4.12). "
        "Twice that median is 6.63 pixels per frame. Three times it is 9.95. Twice the 90th percentile is 8.25. "
        "Optical flow is unreliable where digits overlap. The kept set already drops pixels that are dark in either frame, but overlap boundaries can still invent vectors and fatten the tails. "
        "That bias would have to move the median a long way to change the comparison with an 8-pixel reach. It does not: even a median of 4 would put 2× on the reach and 3× beyond it.",
    )
    add_heading_text(doc, "Q2.2", 2)
    add_body(
        doc,
        "A mirror changes both the appearance and the motion statistics. A mirrored digit is a different shape, and the direction histogram is reflected. "
        "A 90° rotation changes appearance more than a mirror does, and it sends every direction into a different part of Figure 1(a). "
        "A speed-up does not change the digit. It only increases the displacement between the frames the model sees. "
        "The three therefore separate appearance, an unfamiliar direction, and the size of the step. "
        "Constant-velocity motion with elastic bounces is still valid motion after all three, so a model that had learned the motion rule should cope with all of them.",
    )

    add_heading_text(doc, "Table 2. Baseline forecast error", 2)
    rows = [[
        "A",
        learned_row(orig, "MSE1"),
        learned_row(orig, "MSE2"),
        learned_row(orig, "MSE3"),
        learned_row(orig, "MSE4"),
        learned_row(orig, "MSE5"),
        learned_row(orig, "E5"),
        "480,784",
    ]]
    for name, block in (("Blank", blank), ("Copy-last", copy)):
        o = block["original"]
        rows.append([name, fmt(o["MSE1"]), fmt(o["MSE2"]), fmt(o["MSE3"]), fmt(o["MSE4"]), fmt(o["MSE5"]), fmt(o["E5"]), "0"])
    add_table(doc, ["Model", "MSE1", "MSE2", "MSE3", "MSE4", "MSE5", "E(5)", "Parameters"], rows, font=7.5)
    add_caption(doc, "Table 2. Test-set MSE by horizon at 1×, t = 1. Learned entries are mean (seed 0, seed 1).")

    add_figure(
        doc,
        ROOT / "Question2" / "outputs" / "figure2.png",
        "Figure 2. Seed-0 model A on test videos 9000, 9007 and 9012. "
        "Each block is input, prediction, ground truth, five frames each. "
        "Video 9000 is the overlapping pair (a 5 and a 6 merge and separate). "
        "Videos 9007 and 9012 keep two components through frames 1–10.",
        width=5.8,
    )
    add_heading_text(doc, "Q2.3", 2)
    add_body(
        doc,
        "pixel_unshuffle(4) makes each cell of the 16×16 grid stand for 4 original pixels, centre to centre. The 4×4 patch width is not added. "
        "The 1×1 head, the unshuffle and the shuffle do not move information. "
        "Each ConvLSTM layer is a 3×3 convolution with dilation 1, so it reads offsets −1, 0 and +1 and adds a Chebyshev radius of one cell. "
        "In one step the new input passes through layer 1 and then layer 2. Layer 1’s previous hidden state takes that same path. "
        "Layer 2’s previous state passes through layer 2 only, which is shorter. "
        "The longest path is 2 cells, 8 original pixels along a row or a column. "
        "The corner of that neighbourhood is 8√2 ≈ 11.3 pixels. The reach reported here is the axis-aligned one, 8 pixels.",
    )
    add_heading_text(doc, "Q2.4", 2)
    add_body(
        doc,
        "Model A’s E(5) is 0.0223, against 0.0426 for blank and 0.0641 for copy-last. "
        "It is clearly better than both references on ordinary 1× clips. "
        "The error grows with the horizon, from MSE1 = 0.0171 to MSE5 = 0.0263. "
        "That is the closed-loop rollout: each later frame is conditioned on the network’s own previous prediction, so blur accumulates. "
        "In Figure 2 the separated digits in videos 9007 and 9012 stay recognisable and follow the right general direction, while becoming softer. "
        "Video 9000 is the overlap. The 5 and the 6 merge, and the forecast turns them into one soft blob instead of recovering two digits the way the ground truth does. "
        "Overlap is where the model loses which digit is which, not only where it is.",
    )
    add_heading_text(doc, "Q2.5", 2)
    add_body(doc, (ROOT / "Question2" / "predictions" / "Q2.5.txt").read_text(encoding="utf-8").strip())

    add_heading_text(doc, "Table 3. Mirror and rotation", 2)
    add_table(
        doc,
        ["Model", "E(5)", "C_M", "E_M", "C_R", "E_R"],
        [
            [
                "A",
                learned_row(orig, "E5"),
                pair(a["CM"][0], a["CM"][1]),
                pair(a["EM"][0]["E5"], a["EM"][1]["E5"]),
                pair(a["CR"][0], a["CR"][1]),
                pair(a["ER"][0]["E5"], a["ER"][1]["E5"]),
            ],
            ["Blank", fmt(blank["original"]["E5"]), "0", fmt(blank["mirror_error"]["E5"]), "0", fmt(blank["rotation_error"]["E5"])],
            ["Copy-last", fmt(copy["original"]["E5"]), "0", fmt(copy["mirror_error"]["E5"]), "0", fmt(copy["rotation_error"]["E5"])],
        ],
        font=7.5,
    )
    add_caption(
        doc,
        "Table 3. Consistency gaps and transformed error at 1×, t = 1. Blank and copy-last are exactly consistent. Their C_M and C_R are 0.",
    )
    add_figure(
        doc,
        ROOT / "Question2" / "outputs" / "figure3.png",
        "Figure 3. Seed-0 model A, the two test videos with the largest per-video C_M (9642, 9279) and the two with the largest per-video C_R (9168, 9925). "
        "Rows in each panel are X | f(X); T(X) | T(Y); T(X) | f(T(X)); T(X) | T(f(X)). "
        "GIFs: Question2/outputs/gifs/figure3_mirror_0.gif, figure3_mirror_1.gif, figure3_rotation_0.gif, figure3_rotation_1.gif.",
        width=5.8,
    )
    add_figure(
        doc,
        ROOT / "Question2" / "outputs" / "figure4.png",
        "Figure 4. Per-video C_M and C_R of seed-0 model A on the 1,000 test videos. Red lines mark the videos in Figure 3.",
    )
    add_heading_text(doc, "Q2.6", 2)
    add_body(
        doc,
        "C_T = 0 means the forecast commutes with T. It does not mean the forecast is right. "
        "Blank and copy-last commute for a structural reason: a black frame is still black after a mirror or a rotation, and repeating the last frame commutes with a pixel permutation. "
        "Their transformed error equals their original error, 0.0426 and 0.0641, so they are consistent and wrong. "
        "E_T has to sit next to C_T. Otherwise a constant predictor looks perfect.",
    )
    add_heading_text(doc, "Q2.7", 2)
    add_body(
        doc,
        "Video 9168, the largest rotation gap for seed 0 (C_R = 0.00572). "
        "Rows 3 and 4 of that panel share the rotated input. "
        "Row 4 is the original forecast after the same rotation: both digits are still there, blurred, and they keep moving apart. "
        "Row 3 is a fresh forecast from the rotated input: one digit fades within the five horizons and the other smears into a block. "
        "The model’s account of the motion changes. Rotating a forecast it was willing to make on the upright clip does not match the forecast it makes once the clip is rotated. Shown the rotated clip, it loses a digit instead of continuing both trajectories.",
    )
    add_heading_text(doc, "Q2.8", 2)
    add_body(
        doc,
        "Yes. Mean C_R is 0.00293 against mean C_M of 0.00174, and E_R is 0.0239 against E_M of 0.0229 and an original E(5) of 0.0223. "
        "Figure 4 sits further to the right for rotation than for mirror. "
        "From Q2.1, a 90° rotation moves more of the direction histogram than a mirror does. From Q2.2, it also changes the shape of a digit more. "
        "Both of those are about appearance and about which directions were common, not about a different law of motion. "
        "The transformed errors stay far below the blank predictor (0.0426). The model is still forecasting. "
        "The speed test, which changes the step and not the appearance, is the one that reaches the blank error. That is a different failure.",
    )

    def speed_pair(block, k, key):
        return pair(block["speed"][k][0][key], block["speed"][k][1][key])

    add_heading_text(doc, "Table 4. Faster motion", 2)
    add_table(
        doc,
        ["Model", "E(2) at 1×", "E(2) at 2×", "E(2) at 3×", "E(5) at 1×", "E(5) at 2×"],
        [
            ["A", speed_pair(a, "1", "E2"), speed_pair(a, "2", "E2"), speed_pair(a, "3", "E2"), speed_pair(a, "1", "E5"), speed_pair(a, "2", "E5")],
            ["Copy-last", fmt(copy["speed"]["1"]["E2"]), fmt(copy["speed"]["2"]["E2"]), fmt(copy["speed"]["3"]["E2"]), fmt(copy["speed"]["1"]["E5"]), fmt(copy["speed"]["2"]["E5"])],
        ],
        font=7.5,
    )
    add_caption(
        doc,
        "Table 4. Speed comparison on the test set at t = 1. E(5) is not defined at 3×, because only two target frames exist. "
        "Blank’s E(2) is 0.0426, 0.0425 and 0.0426 at 1×, 2× and 3×.",
    )
    add_figure(
        doc,
        ROOT / "Question2" / "outputs" / "figure5.png",
        "Figure 5. Seed-0 model A on test video 9021, chosen so the displayed true frames do not overlap. "
        "Each row is one speed: five input frames, the first two predictions, and the first two true future frames, at 1×, 2× and 3×.",
    )
    add_heading_text(doc, "Q2.9", 2)
    add_body(
        doc,
        "Model A falls apart at 2×, and 3× is worse. E(2) goes from 0.0188 at 1× to 0.0449 at 2× and 0.0523 at 3×. "
        "Blank’s E(2) is about 0.0426, so at 2× the forecast is already no better than predicting black. "
        "This is earlier than the median argument in Q2.5. Twice the median speed is 6.63 pixels, inside a reach of 8, and three times the median is 9.95, outside it. "
        "What that estimate missed is the tail and the width of a digit. Twice the 90th percentile is 8.25 pixels, already past the axis-aligned reach, so a real share of 2× motion cannot be joined to the previous position in one update. A digit is also wider than the one pixel whose flow was measured. "
        "Figure 5, video 9021, shows the failure as smear. At 1× the 1 and the 4 stay sharp enough to follow. At 2× and 3× the forecast is a soft ghost near the last input position, while the true digits have moved on. "
        "A stall would repeat a sharp digit. A duplicate would show two copies. A vanishing digit would go to black. This is the smear: the model no longer has a position to commit to.",
    )
    add_heading_text(doc, "Q2.10", 2)
    add_body(
        doc,
        "No. Model A never trained on k = 2 or k = 3, and 2× motion also sits on the edge of the 8-pixel reach. Either fact can produce Figure 5. "
        "The split is model C, the same architecture trained with k drawn from {1, 2, 3}, against model D, a larger reach trained on the same speeds. "
        "If C repairs 2× and still fails at 3×, then 2× was missing from the training distribution and 3× is beyond the architecture. "
        "If C fails at 2× as well, the reach was already the limit, and more fast clips cannot move information further than the convolutions allow.",
    )
    add_heading_text(doc, "Q2.11", 2)
    add_body(doc, (ROOT / "Question2" / "predictions" / "Q2.11.txt").read_text(encoding="utf-8").strip())
    add_heading_text(doc, "Q2.12", 2)
    add_body(
        doc,
        "Model D is model A’s ConvLSTM with one change. Layer 1 stays a 3×3 with dilation 1. Layer 2 is a 3×3 with dilation 3 and padding 3, so it reads offsets −3, 0 and +3. "
        "Dilation adds no weights, so the parameter count stays 480,784, under the limit of 721,176. There is no attention and no extra layer. "
        "The 9-step rollout, the shapes, the loss and the rest of the training protocol are unchanged. The only protocol change is the one in the model table: k is drawn from {1, 2, 3}. "
        "Reach uses the same rules as Q2.3. One cell is 4 original pixels. Layer 1 adds a radius of 1 cell and layer 2 adds a radius of 3. "
        "The longest path, the new input or layer 1’s previous state through both convolutions, has Chebyshev radius 4 cells, which is 16 original pixels. "
        "The summed offsets are every integer from −4 to 4, so that neighbourhood has no holes. "
        "Layer 2’s own previous state only reaches 12 pixels, which is the shorter path. "
        "16 is strictly larger than model A’s 8. Three times the median speed, 9.95, and three times the 90th percentile, about 12.4, are inside 16 and outside 8.",
    )

    # Table 5 / 6
    headers5 = ["Model", "E(5)", "C_M", "E_M", "C_R", "E_R", "E(2) 1×", "E(2) 2×", "E(2) 3×", "Params", "Train s", "Infer s"]
    add_heading_text(doc, "Table 5. What changes with the data, and with the reach", 2)
    rows5 = []
    for name in ("A", "B", "C", "D", "DB"):
        block = models[name]
        rows5.append([
            "D+B" if name == "DB" else name,
            cell_mean(block, "original", "E5"),
            cell_mean(block, "CM"),
            cell_mean(block, "EM", "E5"),
            cell_mean(block, "CR"),
            cell_mean(block, "ER", "E5"),
            cell_mean(block, "speed", "1", "E2"),
            cell_mean(block, "speed", "2", "E2"),
            cell_mean(block, "speed", "3", "E2"),
            f"{block[0]['parameters']:,}",
            cell_mean(block, "epoch_seconds_mean"),
            cell_mean(block, "inference_seconds"),
        ])
    add_table(doc, headers5, rows5, font=7)
    add_caption(
        doc,
        "Table 5. Test metrics for the learned models. D+B is model D trained with both speed and symmetry augmentation. "
        "Train s is the mean seconds per epoch. Infer s is one forward pass of the 1,000 test videos after the batch is on the Tesla T4.",
    )

    add_heading_text(doc, "Q2.13", 2)
    add_body(
        doc,
        "The mirror and rotation gaps are an appearance failure, in the sense of Q2.2: the motion law is unchanged and the digits are not. "
        f"Model B, which only adds g drawn from G, cuts C_M from {cell_mean(models['A'], 'CM')} to {cell_mean(models['B'], 'CM')} "
        f"and C_R from {cell_mean(models['A'], 'CR')} to {cell_mean(models['B'], 'CR')}. "
        f"Its transformed errors fall into line with its own E(5) of {cell_mean(models['B'], 'original', 'E5')}. "
        f"Model C, which changes speed and not symmetry, leaves C_R at {cell_mean(models['C'], 'CR')}, close to model A. "
        "So the symmetry gap moves when the training distribution is symmetrised, and it does not move when only the speeds change. "
        "The speed failure is unfamiliar motion, not architectural reach. "
        f"Model C has the same 8-pixel reach as A and cuts E(2) at 2× from {cell_mean(models['A'], 'speed', '2', 'E2')} to {cell_mean(models['C'], 'speed', '2', 'E2')} "
        f"and at 3× from {cell_mean(models['A'], 'speed', '3', 'E2')} to {cell_mean(models['C'], 'speed', '3', 'E2')}. "
        "Both of those are better than blank, which sits near 0.0426. "
        f"Model D, with reach 16 and the same speed augmentation, does not beat C: its 3× E(2) is {cell_mean(models['D'], 'speed', '3', 'E2')}. "
        "The conditional in the question does not hold. C is not failing at 3× while D succeeds. "
        "More examples of fast motion were enough, inside the original reach. "
        "That does not mean data can extend a reach. A 3×3 kernel has a zero weight on every pixel outside its offsets, for every value of the weights. "
        "Training changes those weights. It does not change which pixels are read. "
        "D was the test of that limit, and on this data the limit was not what separated the errors.",
    )
    add_heading_text(doc, "Q2.14", 2)
    add_body(
        doc,
        f"Yes. Original E(5) goes from {cell_mean(models['A'], 'original', 'E5')} for A to {cell_mean(models['B'], 'original', 'E5')} for B, "
        f"{cell_mean(models['C'], 'original', 'E5')} for C, {cell_mean(models['D'], 'original', 'E5')} for D and {cell_mean(models['DB'], 'original', 'E5')} for D+B. "
        "B’s cost is small beside the drop in the consistency gaps, and it is worth it. "
        "C’s cost is larger, and it is the cost of no longer collapsing to the blank predictor at 2× and 3×. "
        "That trade is worth it for the question being asked. "
        "D’s further cost over C bought no speed improvement. The dilation was not worth it on these clips. "
        "D+B matches D on the original clips and is the better of those two under mirror and rotation, which is why it is the one wrapped below.",
    )

    add_heading_text(doc, "Table 6. Symmetrised forecasters", 2)
    add_body(
        doc,
        "For any g in G and any T in {M, R}, right-multiplication by T is a bijection of G. "
        "Set h = g T, so g^{−1} = T h^{−1}. Then "
        "f̄(TX) = (1/8) Σ_g g^{−1} f(g T X) = (1/8) Σ_h T(h^{−1} f(h X)) = T(f̄(X)), "
        "because T permutes pixels and therefore passes through the average. "
        "So f̄(TX) = T(f̄(X)), and C_M = C_R = 0 for every underlying f. "
        "The values below are the floating-point remainder.",
    )
    labels = {"B": "sym B", "D": "sym D", "DB": "sym D+B (final)"}
    rows6 = []
    for name in ("B", "D", "DB"):
        block = wrapped[name]
        rows6.append([
            labels[name],
            cell_mean(block, "original", "E5"),
            cell_mean(block, "CM"),
            cell_mean(block, "EM", "E5"),
            cell_mean(block, "CR"),
            cell_mean(block, "ER", "E5"),
            cell_mean(block, "speed", "1", "E2"),
            cell_mean(block, "speed", "2", "E2"),
            cell_mean(block, "speed", "3", "E2"),
            f"{block[0]['parameters']:,}",
            cell_mean(block, "epoch_seconds_mean"),
            cell_mean(block, "inference_seconds"),
        ])
    add_table(doc, headers5, rows6, font=7)
    add_caption(
        doc,
        "Table 6. The same columns as Table 5, for the symmetrised forecasters. "
        "Train s is the underlying learned model. Symmetrisation is not trained. "
        "C_M and C_R are floating-point zeros and are written in scientific notation. "
        "The final model is symmetrised D+B, not symmetrised D.",
    )

    add_figure(
        doc,
        ROOT / "Question2" / "outputs" / "figure6a_DB.png",
        "Figure 6(a). The final model, symmetrised D+B, seed 0, on the same videos as Figure 3. "
        "GIFs: Question2/outputs/gifs/figure6_DB_mirror_0.gif, figure6_DB_mirror_1.gif, "
        "figure6_DB_rotation_0.gif, figure6_DB_rotation_1.gif.",
        width=5.8,
    )
    add_figure(
        doc,
        ROOT / "Question2" / "outputs" / "figure6b_DB.png",
        "Figure 6(b). The same final model on test video 9021, the video in Figure 5, at 1×, 2× and 3×.",
    )
    add_heading_text(doc, "Q2.15", 2)
    add_body(
        doc,
        f"Learned symmetry is model B. Its C_M is {cell_mean(models['B'], 'CM')} and its C_R is {cell_mean(models['B'], 'CR')}, "
        f"against {cell_mean(models['A'], 'CM')} and {cell_mean(models['A'], 'CR')} for A. "
        f"Its E_M and E_R match its E(5) of {cell_mean(models['B'], 'original', 'E5')}. The gap shrinks. It does not become zero. "
        f"Inference stays one forward pass, {cell_mean(models['B'], 'inference_seconds')} s. "
        "Built-in symmetry is the average over G. It is the exact one, because the short argument above does not depend on f. "
        f"Wrapped B has C_M = {cell_mean(wrapped['B'], 'CM')} and C_R = {cell_mean(wrapped['B'], 'CR')}. "
        f"The cost is the eight passes: {cell_mean(wrapped['B'], 'inference_seconds')} s against {cell_mean(models['B'], 'inference_seconds')} s, and "
        f"{cell_mean(wrapped['DB'], 'inference_seconds')} s for the final model against {cell_mean(models['DB'], 'inference_seconds')} s for D+B unwrapped. "
        "Averaging eight forecasts blurs them when they disagree about where a digit is. "
        f"Here they mostly agree. Wrapped B’s E(5) is {cell_mean(wrapped['B'], 'original', 'E5')}, no worse than learned B, "
        f"and wrapped D+B’s E(5) is {cell_mean(wrapped['DB'], 'original', 'E5')}, the same as D+B to three figures. "
        "Built-in symmetry did not hurt accuracy on this test set. It would hurt a model whose eight transformed forecasts pointed different ways.",
    )
    add_heading_text(doc, "Q2.16", 2)
    add_body(
        doc,
        "On video 9168 the disagreement is gone and the forecast is not repaired. "
        "In Figure 3, row 4 kept two blurred digits and row 3 lost one. "
        "In Figure 6(a) those two rows match: both are the same coarse, blocky field. "
        "That match is the group average. It is why C_R is a floating-point zero. "
        "It is not a better account of the bounce. The upright forecast in row 1 is blocky as well, "
        "which lines up with the original E(5) rising from 0.0223 for A to about 0.030 for D+B. "
        "The part of the failure that was a broken symmetry was fixed by the wrapper, and by D+B’s symmetry augmentation before that. "
        "The part that was a lost digit was not turned back into two sharp trajectories. "
        "Figure 6(b) is the same story at speed. The error at 3× is lower than model A’s, and the frames are still smears rather than a 1 and a 4 at the true positions.",
    )
    add_heading_text(doc, "Q2.17", 2)
    add_body(
        doc,
        "The final model has not learned a motion rule that stands without help. "
        "Model C, with A’s 8-pixel reach, cuts E(2) at 2× from 0.0449 to 0.0288 and at 3× from 0.0523 to 0.0336, "
        "so the speed collapse was motion the training clips never showed. "
        "D’s 16-pixel reach does not beat C, so reach was not the binding limit. "
        "Mirror and rotation consistency are exact only because the final forecaster averages all eight symmetries. That is imposed. "
        "Model B learns a smaller gap, not a zero gap. "
        "A skeptic should ask for a change that neither the training distribution nor the wrapper covers. "
        "The test it has not passed is a rotation by an angle other than a multiple of 90°, or a speed above 3×. "
        "Figure 6 is the other warning: the error can fall while the digits turn into blocks.",
    )

    add_heading_text(doc, "AI use", 1)
    add_body(
        doc,
        "I used Cursor to help with understanding the assignment, running notebooks, writing this report, "
        "pushing code to GitHub, and polishing answers. Experiment numbers in the tables come from the saved score files. "
        "Q2.5 and Q2.11 match the prediction files committed before those experiments.",
    )
    add_body(doc, "Example prompts (paraphrased):", bold=True)
    add_body(
        doc,
        "Prompt 1: \"Kaggle failed with ModuleNotFoundError: No module named q2 when I pushed the training script. "
        "What does that mean and how do I fix the import path so the kernel can see my package?\" "
        "Cursor explained that the kernel only sees attached dataset files and helped adjust the entry script so q2 is on the path.",
    )
    add_body(
        doc,
        "Prompt 2: \"Model D uses dilations (1, 3) for reach 16, but its 3× E(2) is slightly worse than Model C in my table. "
        "Why would a larger reach not help here?\" "
        "Cursor walked through unfamiliar motion vs reach, and that dilation 3 can skip nearby cells, so C fixing most of the gap "
        "from seeing fast clips is consistent with the results.",
    )


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # Title page — edit these lines in Word as you like.
    add_body(doc, "Lahore University of Management Sciences", bold=True, size=14)
    add_body(doc, "AI651 — Deep Learning for Space, Time and Graphs", bold=True, size=14)
    add_body(doc, "Assignment 2", bold=True, size=16)
    add_body(doc, "")
    add_body(doc, "Name: Rifa Salman")
    add_body(doc, "Roll number: ")
    add_body(doc, "Fall 2026")
    add_body(doc, "")
    add_body(doc, "GitHub (public): " + REPO)
    add_body(doc, "")
    add_body(
        doc,
        "Question 1 is a design only (not trained). "
        "Question 2 results are from an NVIDIA Tesla T4. "
        "Learned numbers are mean (seed 0, seed 1), three significant figures; "
        "values below 10⁻³ use scientific notation.",
    )
    doc.add_page_break()
    render_q1(doc)
    add_q2(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
    print("bytes", path.stat().st_size)
    # quick sanity: count tables
    from docx import Document as D
    d = D(str(path))
    print("tables", len(d.tables))
